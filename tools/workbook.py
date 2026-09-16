import json,sys
from datetime import datetime, timezone
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

def safe(value):
    if value is None:return None
    text=str(value)
    return "'"+text if text[:1] in ('=','+','-','@') else value

def as_date(value):
    if not value:return None
    if isinstance(value,datetime):return value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value
    try:
        parsed=datetime.fromisoformat(str(value).replace('Z','+00:00'))
        return parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed.tzinfo else parsed
    except ValueError:return value

def canonical_timestamp(value):
    if value is None:return None
    if isinstance(value,datetime):
        parsed=value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
        return parsed.isoformat(timespec='milliseconds').replace('+00:00','Z')
    return str(value)

def style(ws,headers,widths=None):
    ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78');c.alignment=Alignment(vertical='center',wrap_text=True)
    for col in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(col)].width=(widths or {}).get(col,max(16,min(52,len(str(headers[col-1]))+8)))
    ws.row_dimensions[1].height=30
    last=max(2,ws.max_row);tab=Table(displayName='Table_'+''.join(c for c in ws.title if c.isalnum()),ref=f'A1:{get_column_letter(len(headers))}{last}')
    tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)

def add(ws,headers,rows,date_fields=(),widths=None):
    ws.append(headers)
    for row in rows:ws.append([safe(as_date(row.get(h.lower().replace(' ','_'))) if h.lower().replace(' ','_') in date_fields else row.get(h.lower().replace(' ','_'))) for h in headers])
    for col,h in enumerate(headers,1):
        if h.lower().replace(' ','_') in date_fields:
            for cell in ws.iter_cols(min_col=col,max_col=col,min_row=2):
                for c in cell:
                    if isinstance(c.value,datetime):c.number_format='yyyy-mm-dd hh:mm:ss.000'
    style(ws,headers,widths)

def export(src,dest):
 d=json.load(open(src,encoding='utf8'));wb=Workbook();wb.remove(wb.active);wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True
 dash=wb.create_sheet('Dashboard');dash.append(['YouTube Content OS','Value']);dash.append(['Profile',d['profile']]);dash.append(['Status','Local-first; all public content is DRAFT']);dash.append(['Export revision',int(d['revision'])]);dash.append(['Last successful refresh',as_date(d.get('last_successful_refresh'))]);dash.append(['New signals from last refresh',d.get('new_signal_count',0)]);dash.append(['Next human action','Add a verified source, refresh RSS, then curate a DRAFT mission.' if not d['sources'] else 'Review new signals and update DRAFT missions.']);
 budgets=d.get('budgets',[])
 if budgets:
  for b in budgets: dash.append([f"Mission budget: {b['mission_title']}",f"Used {b['used_minutes']} min | Remaining {b['remaining_minutes']} min | Over budget: {b['over_budget']}"])
 else: dash.append(['180-minute budget','No current mission feedback recorded.'])
 style(dash,['YouTube Content OS','Value'],{1:42,2:78});dash['B5'].number_format='yyyy-mm-dd hh:mm:ss.000'
 add(wb.create_sheet('Missions'),['ID','Title','Status','Source ID','Created At'],d['missions'],('created_at',),{1:38,2:60,3:18,4:38,5:27});missions=wb['Missions'];dv_status=DataValidation(type='list',formula1='"DRAFT,READY,ARCHIVED"',allow_blank=False);missions.add_data_validation(dv_status);dv_status.add('C2:C1048576')
 add(wb.create_sheet('Signals'),['ID','Source ID','Video ID','Title','Published At','Views','Fetched At','Human Note'],d['signals'],('published_at','fetched_at'),{1:38,2:38,3:28,4:62,5:27,6:16,7:27,8:54})
 add(wb.create_sheet('Evidence'),['ID','Source ID','Body','Created At'],d['evidence'],('created_at',),{1:38,2:38,3:76,4:27})
 add(wb.create_sheet('Feedback'),['ID','Piece ID','Date','Native retention','ICP signal','Production minutes','Notes'],d['feedback'],('date',),{1:38,2:38,3:18,4:22,5:34,6:22,7:64});fb=wb['Feedback'];dv_retention=DataValidation(type='list',formula1='"UNKNOWN,UP,FLAT,DOWN"',allow_blank=True);fb.add_data_validation(dv_retention);dv_retention.add('D2:D1048576');dv_minutes=DataValidation(type='whole',operator='between',formula1='0',formula2='180',allow_blank=True);fb.add_data_validation(dv_minutes);dv_minutes.add('F2:F1048576');dv_date=DataValidation(type='date',operator='between',formula1='DATE(2000,1,1)',formula2='DATE(2100,12,31)',allow_blank=True);fb.add_data_validation(dv_date);dv_date.add('C2:C1048576')
 cfg=wb.create_sheet('Config');cfg.append(['Setting','Value']);cfg.append(['Profile',d['profile']]);cfg.append(['RSS mode','Enabled']);cfg.append(['API enrichment','Unconfigured: no local API adapter is implemented']);cfg.append(['Channel setup','Each owner supplies a manually verified public channel ID']);style(cfg,['Setting','Value'],{1:28,2:78})
 add(wb.create_sheet('Run Log'),['ID','Kind','Status','Detail','Created At'],d['logs'],('created_at',),{1:38,2:22,3:16,4:70,5:27})
 meta=wb.create_sheet('_Meta');meta.sheet_state='hidden';meta.append(['export_revision',int(d['revision'])]);meta.append(['profile',d['profile']]);meta.append(['schema',3]);wb.save(dest)

def rows(ws,headers):
 result=[]
 for values in ws.iter_rows(min_row=2,values_only=False):
  if all(c.value is None for c in values):continue
  item={}
  for header,cell in zip(headers,values):
   if cell.data_type=='f':raise ValueError('workbook formulas are not importable')
   value=cell.value
   item[header]=canonical_timestamp(value) if header in ('created_at','published_at','fetched_at') else value
  result.append(item)
 return result

def imp(file):
 wb=load_workbook(file,read_only=True,data_only=False)
 if '_Meta' not in wb:raise ValueError('missing workbook schema metadata')
 required={'Missions':['id','title','status','source_id','created_at'],'Signals':['id','source_id','video_id','title','published_at','views','fetched_at','human_note'],'Feedback':['id','piece_id','date','native_retention','icp_signal','production_minutes','notes']}
 if str(wb['_Meta']['B3'].value) != '3':raise ValueError('unsupported workbook schema')
 data={'revision':str(wb['_Meta']['B1'].value),'profile':str(wb['_Meta']['B2'].value)}
 for sheet,headers in required.items():
  if sheet not in wb:raise ValueError('missing workbook sheet '+sheet)
  actual=[str(c.value).lower().replace(' ','_') for c in next(wb[sheet].iter_rows(min_row=1,max_row=1))]
  if actual!=headers:raise ValueError('invalid workbook schema '+sheet)
  data[sheet.lower()]=rows(wb[sheet],headers)
 print(json.dumps(data))
if sys.argv[1]=='export':export(sys.argv[2],sys.argv[3])
else:imp(sys.argv[2])